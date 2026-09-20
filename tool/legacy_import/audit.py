"""Audit + chuẩn hóa CHỈ ĐỌC file Excel "Quản lý tài chính 2026" (KHÔNG ghi file nguồn/DB).

Chạy:  python audit.py --xlsx <file.xlsx> --out <thư mục đầu ra>

Sinh (đều nằm ngoài git, trong thư mục --out):
  audit_stats.json            thống kê cấu trúc/phân phối/đối soát Excel
  real_data_import_rows.csv   mọi dòng nguồn có nghĩa + phân loại + kết quả chuẩn hóa
  real_data_import_review.csv chỉ các dòng cần chủ dự án quyết định
  sim_input.json              giao dịch dự kiến + danh mục cho mô phỏng Financial Engine (Dart)

Không dùng Ghi chú để suy luận ý nghĩa tài chính.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import os
import re
import sys
from collections import Counter, defaultdict

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import normalize as N  # noqa: E402

SHEET = "Ghi chép"
SUMMARY = "Tổng hợp"

CAT_IDS = {
    "Thu nhập": "thu_nhap",
    "Sinh hoạt": "sinh_hoat",
    "Đầu tư": "dau_tu",
    "Tự thưởng": "tu_thuong",
    "Cho đi": "cho_di",
    "Dâng hiến": "dang_hien",
}
STATUS_ORDER = ["Chưa chuẩn bị", "Đã chuẩn bị", "Đã gửi", "Đã dâng"]
OPENING_DATE = "2026-01-01T00:00:00.000"


def redact(v):
    s = "" if v is None else str(v)
    return "<số dài, ẩn>" if re.fullmatch(r"\d{9,}", s.strip()) else v


# --------------------------------------------------------------------------
# Trích xuất thô + audit cấu trúc
# --------------------------------------------------------------------------
def extract(xlsx):
    sha = hashlib.sha256(open(xlsx, "rb").read()).hexdigest()
    wf = openpyxl.load_workbook(xlsx, data_only=False)
    wv = openpyxl.load_workbook(xlsx, data_only=True)
    structure = {"sha256": sha, "sheets": []}
    for ws in wf.worksheets:
        hidden_rows = sorted(r for r, d in ws.row_dimensions.items() if d.hidden)
        nonempty = sum(1 for row in ws.iter_rows() if any(c.value is not None for c in row))
        formulas = sum(1 for row in ws.iter_rows() for c in row if c.data_type == "f")
        structure["sheets"].append({
            "name": ws.title, "state": ws.sheet_state, "dimensions": ws.dimensions,
            "maxRow": ws.max_row, "maxCol": ws.max_column, "nonEmptyRows": nonempty,
            "formulaCells": formulas, "merged": [str(m) for m in ws.merged_cells.ranges],
            "hiddenRows": {"count": len(hidden_rows), "min": hidden_rows[0] if hidden_rows else None,
                           "max": hidden_rows[-1] if hidden_rows else None},
            "hiddenCols": [c for c, d in ws.column_dimensions.items() if d.hidden],
            "autoFilter": ws.auto_filter.ref,
            "comments": [(c.coordinate, c.comment.text) for row in ws.iter_rows() for c in row if c.comment],
        })
    g = wv[SHEET]
    gf = wf[SHEET]
    header = [g.cell(1, c).value for c in range(1, 8)]
    raw, stray = [], []
    for r in range(2, g.max_row + 1):
        v = [g.cell(r, c).value for c in range(1, 8)]
        for c in range(8, g.max_column + 1):
            val = gf.cell(r, c).value
            if val is not None:
                stray.append({"cell": f"{gf.cell(r, c).column_letter}{r}", "value": redact(val)})
        if not any(x is not None and str(x).strip() != "" for x in v[1:]):
            continue
        d = v[5]
        raw.append({"row": r, "cat": v[1], "amt": v[2], "note": v[3], "member": v[4],
                    "date": d.isoformat(timespec="milliseconds") if isinstance(d, dt.datetime) else None,
                    "status": v[6], "hidden": bool(g.row_dimensions[r].hidden),
                    "dateType": type(d).__name__, "amtType": type(v[2]).__name__})
    types = {"date": Counter(x["dateType"] for x in raw), "amount": Counter(x["amtType"] for x in raw),
             "dateFormat": gf["F2"].number_format, "amountFormat": gf["C2"].number_format}
    # Ô Tổng hợp (giá trị cache) để đối soát.
    sv = wv[SUMMARY]
    summary_cells = {}
    for row in sv.iter_rows():
        for c in row:
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                summary_cells[c.coordinate] = c.value
    support = wv["Hỗ trợ kỹ thuật"]
    support_rows = []
    for r in range(2, support.max_row + 1):
        if support.cell(r, 6).value is not None:
            support_rows.append({"row": r, "amount": support.cell(r, 6).value,
                                 "date": str(support.cell(r, 5).value)[:10]})
    return structure, header, raw, stray, types, summary_cells, support_rows


# --------------------------------------------------------------------------
# Mapping sang mô hình app
# --------------------------------------------------------------------------
def proposed_category_id(rec):
    """id danh mục con dự kiến (xác định, không đụng DB)."""
    cat, cls = rec["rawCategory"], rec["cls"]
    if cls == "IMPORT_REVENUE":
        return "thu_nhap"
    if cls == "IMPORT_BUSINESS_EXPENSE":
        return "cpkd_thu_nhap_am"
    if cls == "IMPORT_SPENDING":
        return CAT_IDS[cat]
    if cls == "IMPORT_OTHER_INFLOW":
        return "hoan_" + CAT_IDS[cat]
    if cls in ("IMPORT_SAVINGS_TOPUP", "IMPORT_SAVINGS_WITHDRAW"):
        return "tiet_kiem"
    if cls == "IMPORT_MEMBER_TRANSFER":
        return "chuyen_tien_thanh_vien"
    return None


def category_table(records):
    cats = {
        "thu_nhap": dict(id="thu_nhap", name="Thu nhập", type="income", excludeFromTotals=False, groupKey=None, group="Doanh thu"),
        "cpkd_thu_nhap_am": dict(id="cpkd_thu_nhap_am", name="Chi phí kinh doanh (từ Thu nhập âm)", type="expense", excludeFromTotals=False, groupKey="business_expense", group="Chi phí kinh doanh"),
        "chuyen_tien_thanh_vien": dict(id="chuyen_tien_thanh_vien", name="Chuyển tiền giữa Vợ/Chồng", type="transfer", excludeFromTotals=False, groupKey=None, group="(Chuyển)"),
        "tiet_kiem": dict(id="tiet_kiem", name="Tiết kiệm", type="transfer", excludeFromTotals=False, groupKey=None, group="(Chuyển)"),
        "so_du_dau_ky": dict(id="so_du_dau_ky", name="Số dư đầu kỳ", type="income", excludeFromTotals=True, groupKey=None, group="Khoản thu khác"),
    }
    for name, cid in CAT_IDS.items():
        if name == "Thu nhập":
            continue
        cats[cid] = dict(id=cid, name=name, type="expense", excludeFromTotals=False, groupKey=None, group="Chi tiêu")
        cats["hoan_" + cid] = dict(id="hoan_" + cid, name=f"Hoàn / thu lại · {name}", type="income", excludeFromTotals=True, groupKey=None, group="Khoản thu khác")
    statuses = defaultdict(set)
    for r in records:
        if r["cls"] in N.IMPORT and r["statusProposal"]:
            statuses[proposed_category_id(r)].add(r["statusProposal"])
    for cid, names in statuses.items():
        ordered = [s for s in STATUS_ORDER if s in names]
        cats[cid]["statuses"] = [dict(id=f"st_{cid}_{i}", name=n, sortOrder=i) for i, n in enumerate(ordered)]
    return cats


def to_sim_tx(rec, cats):
    cid = proposed_category_id(rec)
    st = None
    if rec["statusProposal"]:
        for s in cats[cid].get("statuses", []):
            if s["name"] == rec["statusProposal"]:
                st = s["id"]
    m, cm, cls = rec["member"], rec["counterMember"], rec["cls"]
    base = dict(id=f"row-{rec['sourceRow']}", row=rec["sourceRow"], cls=cls, rawCategory=rec["rawCategory"],
                rawSign=1 if rec["rawAmount"] > 0 else -1, categoryId=cid, amountMinor=rec["amountMinor"],
                date=rec["date"], statusId=st, member=m)
    if cls == "IMPORT_REVENUE" or cls == "IMPORT_OTHER_INFLOW":
        base.update(type="income", transferKind=None, sourceKind="external", sourceRefId=None,
                    destinationKind="memberAvailable", destinationRefId=m)
    elif cls in ("IMPORT_BUSINESS_EXPENSE", "IMPORT_SPENDING"):
        base.update(type="expense", transferKind=None, sourceKind="memberAvailable", sourceRefId=m,
                    destinationKind="external", destinationRefId=None)
    elif cls == "IMPORT_SAVINGS_TOPUP":
        base.update(type="transfer", transferKind="savingsTopup", sourceKind="memberAvailable", sourceRefId=m,
                    destinationKind="memberSavingsAsset", destinationRefId=f"savings_unallocated|{m}")
    elif cls == "IMPORT_SAVINGS_WITHDRAW":
        base.update(type="transfer", transferKind="savingsWithdraw", sourceKind="memberSavingsAsset",
                    sourceRefId=f"savings_unallocated|{m}", destinationKind="memberAvailable", destinationRefId=m)
    elif cls == "IMPORT_MEMBER_TRANSFER":
        base.update(type="transfer", transferKind="memberToMember", sourceKind="memberAvailable", sourceRefId=m,
                    destinationKind="memberAvailable", destinationRefId=cm)
    else:  # pragma: no cover
        raise ValueError(cls)
    return base


OPENINGS = [
    ("opening-vo-available", "vo", "memberAvailable", "vo", 627000, "Tổng hợp!P3"),
    ("opening-chong-available", "chong", "memberAvailable", "chong", 1060000, "Tổng hợp!P4"),
    ("opening-vo-savings", "vo", "memberSavingsAsset", "savings_unallocated|vo", 2500000, "Tổng hợp!P5"),
    ("opening-chong-savings", "chong", "memberSavingsAsset", "savings_unallocated|chong", 12000000, "Tổng hợp!P6 (=5.000.000+3.000.000+4.000.000)"),
]


def opening_tx():
    return [dict(id=i, row=0, cls="OPENING_PROPOSAL", rawCategory=None, rawSign=1, categoryId="so_du_dau_ky",
                 amountMinor=a, date=OPENING_DATE, statusId=None, member=m, type="income", transferKind=None,
                 sourceKind="external", sourceRefId=None, destinationKind=k, destinationRefId=ref, evidence=ev)
            for i, m, k, ref, a, ev in OPENINGS]


# --------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)

    structure, header, raw, stray, types, summary_cells, support_rows = extract(a.xlsx)
    recs, mirror_report, dups = N.normalize(raw)
    cats = category_table(recs)

    # ---- CSV toàn bộ dòng + review
    cols = ["sourceRow", "cls", "transactionType", "transferKind", "reportingGroup", "categoryProposal", "member",
            "counterMember", "amountMinor", "date", "statusProposal", "normalizationReason", "flags",
            "rawCategory", "rawAmount", "rawNote", "rawMember", "rawDate", "rawStatus"]

    def rowdict(r):
        d = {k: r.get(k) for k in cols}
        d["flags"] = ";".join(r["flags"])
        return d

    with open(os.path.join(a.out, "real_data_import_rows.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in recs:
            w.writerow(rowdict(r))
    review = [r for r in recs if r["cls"].startswith("REVIEW") or any(x.startswith("EXACT_DUPLICATE") for x in r["flags"])]
    with open(os.path.join(a.out, "real_data_import_review.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in review:
            w.writerow(rowdict(r))

    # ---- Mô phỏng
    sim = [to_sim_tx(r, cats) for r in recs if r["cls"] in N.IMPORT]
    sim_sorted = sorted(sim, key=lambda t: (t["date"], t["row"]))
    payload = {"categories": list(cats.values()), "openings": opening_tx(), "transactions": sim_sorted}
    json.dump(payload, open(os.path.join(a.out, "sim_input.json"), "w", encoding="utf-8"), ensure_ascii=False)
    json.dump(recs, open(os.path.join(a.out, "records.json"), "w", encoding="utf-8"), ensure_ascii=False)

    # ---- Thống kê
    cls_counts = N.class_counts(recs)
    meaningful = len(raw)
    stats = {
        "structure": structure, "header": header, "stray": stray, "types": {
            "date": dict(types["date"]), "amount": dict(types["amount"]),
            "dateFormat": types["dateFormat"], "amountFormat": types["amountFormat"]},
        "supportSheetRows": support_rows, "summaryCells": summary_cells,
        "meaningfulRows": meaningful,
        "transactionLikeRows": sum(1 for r in raw if r["amt"] is not None),
        "classCounts": dict(cls_counts), "classSum": sum(cls_counts.values()),
        "mirrorReport": mirror_report, "duplicateGroups": dups,
    }
    json.dump(stats, open(os.path.join(a.out, "audit_stats.json"), "w", encoding="utf-8"), ensure_ascii=False, default=str, indent=1)
    print("meaningful", meaningful, "classified", sum(cls_counts.values()))
    print(dict(cls_counts))


if __name__ == "__main__":
    main()
